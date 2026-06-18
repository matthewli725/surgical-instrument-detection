"""Generate PPTX from Beamer PDF PNGs + speaker notes extracted from .tex."""

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

PDF_DIR = Path("fdr")
SRC_DIR = PDF_DIR / "source"
SLIDE_DIR = Path("outputs/slides")
OUT = PDF_DIR / "output" / "fdr_presentation.pptx"

# ── Extract speaker notes from .tex ─────────────────────────────────────
TEX_FILE = SRC_DIR / "fdr_presentation.tex"
tex = TEX_FILE.read_text()

# Find all \note{...} blocks using a brace counter
notes: list[str] = []
pos = 0
while True:
    idx = tex.find(r"\note{", pos)
    if idx == -1:
        break
    depth = 1
    i = idx + len(r"\note{")
    start = i
    while i < len(tex) and depth > 0:
        if tex[i] == "{":
            depth += 1
        elif tex[i] == "}":
            depth -= 1
        i += 1
    content = tex[start : i - 1].strip()
    notes.append(content)
    pos = i

n_notes = len(notes)
print(f"Extracted {n_notes} speaker notes from .tex")

# ── Load slide images ──────────────────────────────────────────────────
slide_files = sorted(SLIDE_DIR.glob("slide-*.png"), key=lambda f: int(f.stem.split("-")[1]))
n_slides = len(slide_files)
print(f"Found {n_slides} slide PNGs")

# ── Build PPTX ─────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]  # blank

for i, img_path in enumerate(slide_files):
    slide = prs.slides.add_slide(blank_layout)

    # Add image at full slide size
    pic = slide.shapes.add_picture(
        str(img_path), 0, 0,
        width=prs.slide_width,
        height=prs.slide_height,
    )

    # Add speaker note
    if i < n_notes:
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = notes[i]

out_path = str(OUT)
prs.save(out_path)
print(f"Saved {out_path}  ({n_slides} slides)")
